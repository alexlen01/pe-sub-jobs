#!/usr/bin/env python3
r"""The generated LP DB Export must obey the rules it is a sample OF.

lp_db_generate.py randomizes a sample; it must not corrupt one. Two halves are asserted here:

  * what the WORKBOOK says is internally true - every classification is one the platform knows,
    every "Included Investors (Rated)" LP carries an agency rating and no other LP does, no NR/N/A token is
    ever written, the UBS rate and limit are the ones the Borrowing Base Criteria Matrix states for
    that classification, the agent rate and limit are the ones the Agent Advance Rate Schedule
    states for its CLASSIFICATION, and both borrowing bases re-derive from the row's own inputs;

  * what the chaos monkey does to it is RECOVERABLE - it only ever blanks a cell with a documented
    fallback or respells a value the extract's reference lists resolve back, so lp_db_extract.py
    turns the degraded workbook back into the canonical values the clean one held.

The reconciliation here is deliberately an independent re-implementation of BbCalculationService's
computeOne rather than a call into the generator's own verify_reconciliation: a check that shares
its subject's arithmetic cannot fail when that arithmetic is wrong.

Run:  python pe-sub-jobs/scripts/tests/test_generator_rules.py
"""
from __future__ import annotations

import contextlib
import csv
import importlib.util
import io
import shutil
import sys
import tempfile
import unittest
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parent.parent
REFERENCE_DIR = SCRIPTS_DIR.parent / "data" / "reference"

SAMPLE_ROWS = 700       # enough for every classification and both funded sides to appear
MONEY_TOL = 0.011       # the workbook carries cents; a cent of rounding either way is not a break
RATE_TOL = 1e-9

_TMP = Path(tempfile.mkdtemp(prefix="lp_db_generate_rules_"))
_SAMPLES: dict[tuple, Path] = {}


def _load(name: str, alias: str | None = None):
    """Import one of the scripts under its own name (or an alias, for a second instance)."""
    spec = importlib.util.spec_from_file_location(alias or name, SCRIPTS_DIR / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod          # @dataclass resolves annotations through sys.modules
    spec.loader.exec_module(mod)
    return mod


def generate(chaos: bool, chaos_seed: int | None = None, tag: str = "") -> tuple[object, list[dict]]:
    """(the generator module, its rows as header->value dicts). Cached: one run per distinct
    configuration, because generating is the slow part of this file."""
    key = (chaos, chaos_seed, tag)
    gen = _load("lp_db_generate", f"lp_db_generate_{len(_SAMPLES)}_{tag}")
    gen.EXPORT_OUT = _TMP / f"export_{chaos}_{chaos_seed}_{tag}.xlsx"
    gen.TARGET_ROWS = SAMPLE_ROWS
    gen.CHAOS_ENABLED = chaos
    if chaos_seed is not None:
        gen.CHAOS_SEED = chaos_seed
    if key not in _SAMPLES:
        with contextlib.redirect_stdout(io.StringIO()):
            if gen.main() != 0:
                raise AssertionError("the generator refused to write a sample")
        _SAMPLES[key] = gen.EXPORT_OUT
    return gen, read_workbook(_SAMPLES[key])


def read_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    out = [dict(zip(header, r)) for r in rows]
    wb.close()
    return out


def num(v) -> float:
    return float(v)


def blank(v) -> bool:
    return v is None or str(v).strip() == ""


def seed_frac(v) -> float:
    """A rate or limit as the seed CSV renders it ('90%', '12.5%') back to a fraction."""
    return float(str(v).strip().rstrip("%")) / 100


def tranche_haircut(gen, rate: float, limit: float) -> tuple[float, float]:
    """The uncommitted sleeve's price for a committed (rate, limit), derived here rather than taken
    from the generator: a rung down AGENT_RATE_LADDER and half the concentration cap. The rung is
    matched by proximity so the seed CSV's '90%' and the ladder's 0.90 cannot miss each other on a
    float comparison."""
    ladder = gen.AGENT_RATE_LADDER
    rung = min(range(len(ladder)), key=lambda k: abs(ladder[k] - rate))
    return ladder[min(rung + 1, len(ladder) - 1)], limit * gen.UNCOMMITTED_CL_FACTOR


def fac(row: dict) -> tuple[str, str]:
    """The facility a workbook row belongs to: its (AccountID, FndName) pair. One account can carry
    two funds and one fund can be reported under two accounts, so neither half groups a borrowing
    base on its own."""
    return row["AccountID"], row["FndName"]


def stated_ratings(row: dict) -> list[str]:
    return [str(row[c]).strip() for c in ("S&P", "Moody'S", "Fitch") if not blank(row[c])]


def tearDownModule():
    shutil.rmtree(_TMP, ignore_errors=True)


# ── what the workbook says ──────────────────────────────────────────────────────────────────
class CleanSampleTest(unittest.TestCase):
    """Chaos off: the values as the generator decided them, before any spelling is degraded."""

    @classmethod
    def setUpClass(cls):
        cls.gen, cls.rows = generate(chaos=False)
        cls.agent_schedule = {c: (rate, limit) for c, rate, limit, _ in cls.gen.AGENT_CATEGORIES}

    def test_every_ubs_classification_is_one_of_the_nine(self):
        got = {r["UBS LP Classification"] for r in self.rows}
        self.assertTrue(got <= set(self.gen.UBS_CLASSES),
                        f"classes the platform does not know: {got - set(self.gen.UBS_CLASSES)}")

    def test_every_agent_classification_is_one_the_schedule_states(self):
        """The Agent Advance Rate Schedule's CLASSIFICATION column is the whole vocabulary."""
        schedule = {r[0] for r in self.gen._reference_rows("agent_rate_map.csv")[1:] if r}
        got = {r["Agent LP Classification"] for r in self.rows}
        self.assertEqual(set(self.agent_schedule), schedule,
                         "the generator's agent categories have drifted from agent_rate_map.csv")
        self.assertTrue(got <= schedule, f"not on the schedule: {got - schedule}")

    def test_ratings_are_exactly_the_rated_included_population(self):
        """A rating is the EVIDENCE for the bucket, so the two can never disagree: every Rated
        Included LP carries at least one agency, and no other LP carries any."""
        for i, r in enumerate(self.rows, 2):
            with self.subTest(row=i):
                self.assertEqual(r["Agent LP Classification"] == "Included Investors (Rated)",
                                 bool(stated_ratings(r)))

    def test_the_two_taxonomies_agree_at_both_ends(self):
        """Included Investors (Rated) <-> Rated Investor and Excluded Investors <-> Excluded. The pairs the
        business called out: an LP cannot be Excluded to one side and Included Investors (Rated) to the other."""
        for i, r in enumerate(self.rows, 2):
            with self.subTest(row=i):
                agent, ubs = r["Agent LP Classification"], r["UBS LP Classification"]
                self.assertEqual(agent == "Included Investors (Rated)", ubs == "Rated Investor")
                self.assertEqual(agent == "Excluded Investors", ubs == "Excluded")

    def test_no_not_rated_token_is_ever_written(self):
        """NR / N/R / N/A never appear in the real export - an agency that does not rate the LP
        leaves the cell EMPTY - and to pe-sub-api a token reads as a rating it cannot band."""
        ext = _load("lp_db_extract")
        for i, r in enumerate(self.rows, 2):
            for v in stated_ratings(r):
                with self.subTest(row=i, value=v):
                    self.assertNotIn("".join(ch for ch in v.lower() if ch.isalpha()),
                                     ext.NOT_RATED_TOKENS)

    def test_every_notch_is_on_its_own_agencys_scale(self):
        """S&P and Fitch write AA-/BBB+ where Moody's writes Aa3/Baa1. A notch in the wrong column
        matches no band and clamps the LP to the BBB floor, quietly moving its advance rate."""
        for agency, col in (("sp", "S&P"), ("moodys", "Moody'S"), ("fitch", "Fitch")):
            ladder = {n for band in self.gen.RATING_BANDS[agency].values() for n in band}
            ladder |= set(self.gen.SUB_IG_NOTCHES[agency])
            for i, r in enumerate(self.rows, 2):
                if not blank(r[col]):
                    with self.subTest(row=i, agency=agency, notch=r[col]):
                        self.assertIn(str(r[col]).strip(), ladder)

    def test_ubs_rate_and_limit_are_the_matrix_values(self):
        """The export STATES both, and pe-sub-api prefers a stated value over its own default - so a
        rate that is not the matrix's is not a dirty rate, it is a different price."""
        for i, r in enumerate(self.rows, 2):
            band = self.gen.rating_band(r["S&P"], r["Moody'S"], r["Fitch"])
            funded = num(r["Called Capital"]) / num(r["Capital Commitments"])
            rate, limit = self.gen.bb_criteria(r["UBS LP Classification"], band, funded)
            with self.subTest(row=i, cls=r["UBS LP Classification"], band=band):
                self.assertAlmostEqual(rate, num(r["UBS Advance Rate"]), delta=RATE_TOL)
                self.assertAlmostEqual(limit, num(r["UBS Concentration Limit"]), delta=RATE_TOL)

    def test_the_funded_split_is_actually_exercised(self):
        """The matrix prices a Rated Investor differently below and above 40% called. If every
        sample row sat on one side, the test above would be asserting half a rule."""
        sides = Counter(num(r["Called Capital"]) / num(r["Capital Commitments"])
                        >= self.gen.FUNDED_THRESHOLD_PCT for r in self.rows)
        self.assertTrue(sides[True] and sides[False], f"only one funded side present: {sides}")

    def test_agent_rate_and_limit_are_the_schedule_values(self):
        """The Agent Advance Rate Schedule prices the row's CLASSIFICATION - on an uncommitted
        tranche, one rung down the ladder and against half the concentration cap, because that
        sleeve is a separate facility under the Credit Agreement and the agent certifies it
        separately."""
        for i, r in enumerate(self.rows, 2):
            rate, limit = self.agent_schedule[r["Agent LP Classification"]]
            if self.gen.tranche_of(r["FndName"]) == "Uncommitted":
                rate, limit = tranche_haircut(self.gen, rate, limit)
            with self.subTest(row=i, cls=r["Agent LP Classification"], fund=r["FndName"]):
                self.assertAlmostEqual(rate, num(r["Agent Advance Rate"]), delta=RATE_TOL)
                self.assertAlmostEqual(limit, num(r["Agent Concentration Limit"]), delta=RATE_TOL)

    def test_the_cash_columns_add_up(self):
        for i, r in enumerate(self.rows, 2):
            with self.subTest(row=i):
                self.assertAlmostEqual(num(r["Capital Commitments"]) - num(r["Uncalled Capital"]),
                                       num(r["Called Capital"]), delta=MONEY_TOL)
                self.assertAlmostEqual(num(r["Called Capital"]) / num(r["Capital Commitments"]),
                                       num(r["% of LP Called"]), delta=1e-8)

    def test_both_borrowing_bases_re_derive_from_the_rows_own_inputs(self):
        """An independent re-implementation of BbCalculationService.computeOne, which is the engine
        that will re-run this file: eligible uncalled is capped at the LP's concentration limit of
        the facility's TOTAL uncalled, the base advances only that, and an Excluded LP's whole
        uncalled is excess on the UBS side while the agent side reports no excess at all."""
        total_uncalled = defaultdict(float)
        for r in self.rows:
            total_uncalled[fac(r)] += num(r["Uncalled Capital"])

        for i, r in enumerate(self.rows, 2):
            with self.subTest(row=i):
                total = total_uncalled[fac(r)]
                uncalled = num(r["Uncalled Capital"])
                excluded = r["UBS LP Classification"] == "Excluded"

                ubs_eligible = 0.0 if excluded else min(uncalled, num(r["UBS Concentration Limit"]) * total)
                self.assertAlmostEqual(uncalled - ubs_eligible,
                                       num(r["UBS Excess Concentration"]), delta=MONEY_TOL)
                self.assertAlmostEqual(ubs_eligible * num(r["UBS Advance Rate"]),
                                       num(r["UBS Borrowing Base"]), delta=MONEY_TOL)

                agent_limit = num(r["Agent Concentration Limit"])
                agent_eligible = uncalled if agent_limit <= 0 else min(uncalled, agent_limit * total)
                agent_excess = 0.0 if excluded else uncalled - agent_eligible
                self.assertAlmostEqual(agent_excess,
                                       num(r["Agent Excess Concentration"]), delta=MONEY_TOL)
                self.assertAlmostEqual(0.0 if excluded else (uncalled - agent_excess) * num(r["Agent Advance Rate"]),
                                       num(r["Agent Borrowing Base"]), delta=MONEY_TOL)

    def test_the_percentage_shares_are_shares_of_their_facility(self):
        totals = defaultdict(lambda: [0.0, 0.0])
        for r in self.rows:
            t = totals[fac(r)]
            t[0] += num(r["Capital Commitments"])
            t[1] += num(r["Uncalled Capital"])
        for i, r in enumerate(self.rows, 2):
            commit_total, uncalled_total = totals[fac(r)]
            with self.subTest(row=i):
                self.assertAlmostEqual(num(r["Capital Commitments"]) / commit_total,
                                       num(r["% of Capital Commitments"]), delta=1e-8)
                self.assertAlmostEqual(num(r["Uncalled Capital"]) / uncalled_total,
                                       num(r["% of Uncalled Capital"]), delta=1e-8)


# ── the two sleeves of one Credit Agreement ────────────────────────────────────────
class TrancheTest(unittest.TestCase):
    """"<Fund> (Committed)" and "<Fund> (Uncommitted)" are the same fund financed twice under one
    Credit Agreement: a committed line and an uncommitted sleeve drawn only with lender consent.
    The real LP DB export carries both, with the same LPs and the same money under two AccountIDs,
    and the business agreed on Day 1 that the platform tracks each as its own record - so the
    sample has to present them the way the source does, or the ingestion it exercises never meets
    the case at all.
    """

    @classmethod
    def setUpClass(cls):
        cls.gen, cls.rows = generate(chaos=False)
        cls.by_fac = defaultdict(list)
        for r in cls.rows:
            cls.by_fac[fac(r)].append(r)
        # Only the pairs the draw actually reached: a small sample does not place a position on
        # every facility the report lists, and an empty facility asserts nothing either way.
        cls.pairs = [(lead, sib)
                     for lead, sibs in cls.gen.tranche_groups(cls.gen.load_facilities())[0].items()
                     for sib in sibs if cls.by_fac[lead]]

    def test_the_sample_actually_contains_tranche_pairs(self):
        """Every assertion below is vacuous without one, and the pairs come from the summary report,
        where a reissued facility list could quietly drop them."""
        self.assertTrue(self.pairs, "no tranche pair reached the sample - nothing here is tested")
        for lead, sib in self.pairs:
            with self.subTest(lead=lead, sibling=sib):
                self.assertTrue(self.by_fac[sib], f"{sib} mirrors {lead} but drew no positions")

    def test_the_sleeves_are_separate_facilities(self):
        """Separate records, not a duplicate to be collapsed: different AccountID, different fund
        name, and a name that still says which sleeve it is."""
        for lead, sib in self.pairs:
            a, b = self.by_fac[lead][0], self.by_fac[sib][0]
            with self.subTest(lead=lead, sibling=sib):
                self.assertNotEqual(a["AccountID"], b["AccountID"])
                self.assertNotEqual(a["FndName"], b["FndName"])
                self.assertEqual(self.gen.tranche_base_name(a["FndName"]),
                                 self.gen.tranche_base_name(b["FndName"]))
                self.assertEqual({"Committed", "Uncommitted"},
                                 {self.gen.tranche_of(a["FndName"]), self.gen.tranche_of(b["FndName"])})

    def test_both_sleeves_carry_the_identical_lp_roster_and_money(self):
        """One set of LPs pledging one pool of collateral under one Credit Agreement: the same LPs,
        the same COUNT of them, and the same commitment, called and uncalled capital on each side.
        This is the property the export is recognised by - the pair that looks duplicated because
        only the account number differs."""
        for lead, sib in self.pairs:
            L, S = self.by_fac[lead], self.by_fac[sib]
            with self.subTest(lead=lead, sibling=sib):
                self.assertEqual(len(L), len(S), "the sleeves carry a different number of LPs")
                for a, b in zip(L, S):
                    self.assertEqual(a["Investor Name"], b["Investor Name"])
                    for col in ("Capital Commitments", "Called Capital", "Uncalled Capital"):
                        self.assertAlmostEqual(num(a[col]), num(b[col]), delta=MONEY_TOL, msg=col)

    def test_the_uncommitted_sleeve_is_priced_and_certified_lower(self):
        """The reason the split has to survive ingestion: identical collateral, two borrowing bases.
        Every LP is haircut a rung and tested against half the cap on the uncommitted sleeve, so its
        agent borrowing base comes out strictly smaller - except an Excluded LP, which contributes
        nothing on either side."""
        for lead, sib in self.pairs:
            L, S = self.by_fac[lead], self.by_fac[sib]
            with self.subTest(lead=lead, sibling=sib):
                for a, b in zip(L, S):
                    if num(a["Agent Advance Rate"]) == 0:          # Excluded Investors
                        self.assertEqual(0, num(b["Agent Borrowing Base"]))
                        continue
                    self.assertLess(num(b["Agent Advance Rate"]), num(a["Agent Advance Rate"]))
                    self.assertLess(num(b["Agent Concentration Limit"]),
                                    num(a["Agent Concentration Limit"]))
                    self.assertLess(num(b["Agent Borrowing Base"]), num(a["Agent Borrowing Base"]))
                self.assertLess(sum(num(r["Agent Borrowing Base"]) for r in S),
                                sum(num(r["Agent Borrowing Base"]) for r in L))

    def test_the_ubs_columns_do_not_move_with_the_sleeve(self):
        """UBS's matrix prices the LP's classification, not the tranche - and the extract
        consolidates UBSAR/UBSCL by recency into LP Master's BANK-WIDE default. A haircut written
        there would escape the facility it belongs to and re-price the LP everywhere."""
        for lead, sib in self.pairs:
            for a, b in zip(self.by_fac[lead], self.by_fac[sib]):
                with self.subTest(lead=lead, sibling=sib, investor=a["Investor Name"]):
                    self.assertEqual(a["UBS LP Classification"], b["UBS LP Classification"])
                    self.assertAlmostEqual(num(a["UBS Advance Rate"]), num(b["UBS Advance Rate"]),
                                           delta=RATE_TOL)
                    self.assertAlmostEqual(num(a["UBS Concentration Limit"]),
                                           num(b["UBS Concentration Limit"]), delta=RATE_TOL)


# ── what the chaos monkey is allowed to do to it ────────────────────────────────────────────
class ChaosBoundsTest(unittest.TestCase):
    """Degradation is a spelling change or a blank with a fallback - never a change of meaning."""

    @classmethod
    def setUpClass(cls):
        cls.gen, cls.clean = generate(chaos=False)
        _, cls.dirty = generate(chaos=True)

    def test_chaos_leaves_the_base_data_alone(self):
        self.assertEqual(len(self.clean), len(self.dirty))

    def test_the_same_chaos_seed_reproduces_the_same_workbook(self):
        """The sample is a fixture: an investigation into a row it produced has to be repeatable."""
        _, again = generate(chaos=True, tag="repeat")
        self.assertEqual(self.dirty, again)

    def test_a_different_chaos_seed_degrades_differently(self):
        """Guards the reproducibility test above against passing because chaos did nothing."""
        _, other = generate(chaos=True, chaos_seed=987654321, tag="other")
        self.assertNotEqual(self.dirty, other)

    def test_the_sacred_columns_are_untouched(self):
        """The cash, the legal LPA figures, the facility join keys, the whole UBS credit family
        (classification, advance rate, concentration limit, borrowing base and the excess that nets
        it) and every column the borrowing bases reconcile from. A dirty spelling of a name is
        still that name; a dirty commitment is a different commitment.

        The internal-key -> header map is taken from the generator's own two column lists rather
        than restated here, so a column added to CHAOS_SACRED is covered the moment it is added."""
        # Keyed positionally: openpyxl normalizes the CRLF inside the "LP Size" header on read, so
        # the workbook's own header row is what the row dicts are keyed by, not SRC_HEADERS.
        headers = list(self.clean[0])
        header_of = {col: headers[i] for i, col in enumerate(self.gen.SRC_COLS)}
        for col in self.gen.CHAOS_SACRED:
            header = header_of[col]
            with self.subTest(column=header):
                self.assertEqual([r[header] for r in self.clean], [r[header] for r in self.dirty])

    def test_only_cells_with_a_fallback_are_ever_blanked(self):
        """A blank is legal exactly where something else resolves the value that was there: the
        Agent Advance Rate Schedule for AgentAR/AgentCL, and an earlier submission of the same LP
        for the size basis. The UBS rate and limit are NOT among them - nothing downstream restores
        the seeded value, so an emptied cell would re-price the LP rather than leave a gap."""
        blankable = {"AgentAR": "Agent Advance Rate", "AgentCL": "Agent Concentration Limit",
                     "LpSizeCriteria": "LP Size Criteria"}
        allowed = {blankable[c] for c in self.gen.CHAOS_BLANKABLE}
        for i, (clean, dirty) in enumerate(zip(self.clean, self.dirty), 2):
            for header, value in clean.items():
                if not blank(value) and blank(dirty[header]):
                    with self.subTest(row=i, column=header):
                        self.assertIn(header, allowed, "blanked a cell with no documented fallback")


# ── and that the extract gets it all back ───────────────────────────────────────────────────
class ChaosRoundTripTest(unittest.TestCase):
    """The degraded workbook, read by lp_db_extract, yields the canonical values the clean one held.
    This is the property that makes the chaos monkey a randomizer rather than a corrupter."""

    @classmethod
    def setUpClass(cls):
        cls.gen, cls.clean = generate(chaos=False)
        _, cls.dirty = generate(chaos=True)

        tmp = _TMP / "roundtrip"
        (tmp / "out").mkdir(parents=True, exist_ok=True)
        report = openpyxl.Workbook()
        rs = report.active
        rs.append(["Agent", "Borrower", "AccountNumber", "LoanAmount", None,
                   "MaturityDate", "FacilityStatus", "FacilityStatusDate"])
        rs.append(["Bank of America", None, None, None, None, None, None, None])
        for acct, fnd in dict.fromkeys((r["AccountID"], r["FndName"]) for r in cls.dirty):
            rs.append([None, fnd, acct, "100000000", None, "2030-01-01", "Active", "2026-01-01"])
        rs.append([None, None, None, "AccessTotalsLoanAmount:", None, None, None, None])
        report.save(tmp / "AgentBankSummaryRpt.xlsx")

        cls.ext = _load("lp_db_extract")
        cls.ext.EXPORT_FILE = _SAMPLES[(True, None, "")]
        cls.ext.AGENT_BANK_SUMMARY_FILE = tmp / "AgentBankSummaryRpt.xlsx"
        cls.ext.OUT_DIR = tmp / "out"
        cls.ext.REFERENCE_DIR = REFERENCE_DIR
        with contextlib.redirect_stdout(io.StringIO()):
            if cls.ext.main() != 0:
                raise AssertionError("the extract refused the degraded workbook")
        cls.seeds = list(csv.DictReader((tmp / "out" / "lp_facility_seeds.csv").open(encoding="utf-8")))
        cls.master = list(csv.DictReader((tmp / "out" / "lp_master.csv").open(encoding="utf-8")))

    def test_no_record_is_lost(self):
        self.assertEqual(len(self.dirty), len(self.seeds))

    def test_every_agent_classification_comes_back_canonical(self):
        schedule = {c for c, *_ in self.gen.AGENT_CATEGORIES}
        got = Counter(s["agent_lp_category"] for s in self.seeds)
        self.assertEqual([], [c for c in got if c not in schedule],
                         f"did not resolve back to the schedule: {dict(got)}")

    def test_every_ubs_classification_comes_back_canonical(self):
        got = Counter(s["ubs_lp_category"] for s in self.seeds)
        self.assertEqual([], [c for c in got if c not in self.gen.UBS_CLASSES],
                         f"did not resolve back to the nine: {dict(got)}")

    def test_the_classification_mix_survives_the_round_trip(self):
        """Resolving is not enough - each row has to come back as the class it left as, or the
        degradation moved LPs between buckets rather than respelling their labels."""
        self.assertEqual([r["UBS LP Classification"] for r in self.clean],
                         [s["ubs_lp_category"] for s in self.seeds])
        self.assertEqual([r["Agent LP Classification"] for r in self.clean],
                         [s["agent_lp_category"] for s in self.seeds])

    def test_every_rating_comes_back_on_its_agencys_own_scale(self):
        """'A minus' and 'Baa 1' are how an analyst types A- and Baa1. pe-sub-api matches the notch
        exactly and clamps what it cannot find to the BBB floor, so normalizing is what stops a
        spelling difference from becoming a pricing difference."""
        for agency, col in (("sp", "sp_rating"), ("moodys", "moodys_rating"), ("fitch", "fitch_rating")):
            ladder = {n for band in self.gen.RATING_BANDS[agency].values() for n in band}
            ladder |= set(self.gen.SUB_IG_NOTCHES[agency])
            for s in self.seeds:
                if s[col].strip():
                    with self.subTest(agency=agency, notch=s[col]):
                        self.assertIn(s[col].strip(), ladder)

    def test_the_rates_and_limits_a_blank_cell_lost_are_resolved_back(self):
        """AgentAR / AgentCL are blanked by chaos and recovered from the Agent Advance Rate
        Schedule, so the seed must carry the schedule's value for the row's classification.

        An uncommitted tranche is the exception that proves the rule: its cells state that sleeve's
        own haircut price, which the schedule could not resolve a blank back to, so the chaos monkey
        never empties them and the seed still comes back at the price the workbook stated."""
        schedule = {c: (rate, limit) for c, rate, limit, _ in self.gen.AGENT_CATEGORIES}
        for s, clean in zip(self.seeds, self.clean):
            rate, limit = schedule[s["agent_lp_category"]]
            if self.gen.tranche_of(clean["FndName"]) == "Uncommitted":
                rate, limit = tranche_haircut(self.gen, rate, limit)
            with self.subTest(investor=s["investor_name"], cls=s["agent_lp_category"]):
                self.assertAlmostEqual(rate, seed_frac(s["agent_advance_rate"]), delta=RATE_TOL)
                self.assertAlmostEqual(limit, seed_frac(s["agent_concentration_limit"]),
                                       delta=RATE_TOL)

    def test_name_drift_does_not_split_an_lp_into_several_profiles(self):
        """lp_master is one row per LP. A name that drifted row to row would fan one LP out into
        several partial profiles, which is a corrupted sample, not a degraded one."""
        self.assertEqual(len({r["Investor Name"] for r in self.clean}), len(self.master))
        dupes = {n: c for n, c in Counter(m["investor_name"] for m in self.master).items() if c > 1}
        self.assertEqual({}, dupes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
